from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


_HEADINGS = [
    "摘要", "问题重述", "问题分析", "模型假设", "符号说明",
    "模型的建立与求解", "灵敏度分析", "模型评价与推广", "参考文献",
    "AI 工具使用声明", "附录：支撑文件/材料清单", "附录：程序代码",
]


def _set_style_font(style, east_asia: str, latin: str, size: float, bold: bool = False) -> None:
    style.font.name = latin
    style.font.size = Pt(size)
    style.font.bold = bold
    style.element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    style.element.rPr.rFonts.set(qn("w:ascii"), latin)
    style.element.rPr.rFonts.set(qn("w:hAnsi"), latin)


def _set_run_font(run, east_asia: str, latin: str = "Times New Roman") -> None:
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)
    section.top_margin = Cm(2.10)
    section.bottom_margin = Cm(2.20)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    _add_page_number(section.footer.paragraphs[0])


def _configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    _set_style_font(normal, "宋体", "Times New Roman", 12)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = 1.32
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    _set_style_font(title, "黑体", "Times New Roman", 22, bold=True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True

    heading_1 = styles["Heading 1"]
    _set_style_font(heading_1, "黑体", "Times New Roman", 16, bold=True)
    heading_1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading_1.paragraph_format.space_before = Pt(14)
    heading_1.paragraph_format.space_after = Pt(7)
    heading_1.paragraph_format.keep_with_next = True
    heading_1.paragraph_format.keep_together = True

    heading_2 = styles["Heading 2"]
    _set_style_font(heading_2, "黑体", "Times New Roman", 14, bold=True)
    heading_2.paragraph_format.space_before = Pt(10)
    heading_2.paragraph_format.space_after = Pt(4)
    heading_2.paragraph_format.keep_with_next = True

    heading_3 = styles["Heading 3"]
    _set_style_font(heading_3, "黑体", "Times New Roman", 12, bold=True)
    heading_3.paragraph_format.space_before = Pt(8)
    heading_3.paragraph_format.space_after = Pt(3)
    heading_3.paragraph_format.keep_with_next = True

    for name in ("图题", "表题"):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        _set_style_font(style, "宋体", "Times New Roman", 10.5)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = name == "表题"

    note = styles.add_style("表注", WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(note, "宋体", "Times New Roman", 9)
    note.paragraph_format.first_line_indent = Pt(0)
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(3)


def _add_placeholder(doc: Document, text: str = "（占位段落，赛时替换为正文。）") -> None:
    doc.add_paragraph(text)


def build_template(out_path: str) -> None:
    doc = Document()
    doc.core_properties.title = "全国大学生数学建模竞赛论文模板"
    doc.core_properties.subject = "cumcm-toolkit B226 竞赛型排版"
    _configure_page(doc)
    _configure_styles(doc)

    doc.add_heading("论文标题", level=0)

    doc.add_heading("摘要", level=1)
    _add_placeholder(doc, "摘要应在一段内依次交代问题、方法、关键结果与创新点，避免只复述题目。")
    keywords = doc.add_paragraph()
    keywords.paragraph_format.first_line_indent = Pt(0)
    label = keywords.add_run("关键词：")
    label.bold = True
    _set_run_font(label, "黑体")
    keywords.add_run("关键词1；关键词2；关键词3；关键词4")

    for heading in ("问题重述", "问题分析"):
        doc.add_heading(heading, level=1)
        _add_placeholder(doc)

    doc.add_heading("模型假设", level=1)
    doc.add_paragraph("假设 1：写明假设内容及其合理性。", style="List Number")

    doc.add_heading("符号说明", level=1)
    doc.add_paragraph("表 1  主要符号说明", style="表题")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("符号", "含义", "单位")):
        cell.text = value
    for cell, value in zip(table.rows[1].cells, ("Z", "目标函数", "按题意")):
        cell.text = value
    for cell, value in zip(table.rows[2].cells, ("x_i", "第 i 个决策变量", "按题意")):
        cell.text = value
    doc.add_paragraph("符号首次出现时仍需在正文中解释。", style="表注")

    doc.add_heading("模型的建立与求解", level=1)
    doc.add_heading("问题一", level=2)
    # 这四个小节承载 analysis/modeling-workflow.md 第 6 节契约的六项内容：
    # 「问题分析」在顶层同名章节，「求解结果」与「结果解释」合并进「结果分析」。
    for heading, prompt in (
        ("模型建立", "定义变量、参数、目标或待求量、约束和核心方程，并解释推导。"),
        ("算法设计与求解", "说明基线、主方法、算法步骤、停止条件和结果来源。"),
        ("结果分析", "报告方案、目标值、约束可行性和相对基线的改善，并解释结果为何合理、适用边界在哪。"),
        ("模型验证", "按任务类型给出收敛、多种子、独立算法、残差或敏感性证据。"),
    ):
        doc.add_heading(heading, level=3)
        _add_placeholder(doc, prompt)
    doc.add_paragraph("图 1  结果图标题", style="图题")

    doc.add_heading("灵敏度分析", level=1)
    _add_placeholder(doc)

    doc.add_heading("模型评价与推广", level=1)
    doc.add_heading("模型的优点", level=2)
    _add_placeholder(doc)
    doc.add_heading("模型的不足与推广", level=2)
    _add_placeholder(doc)

    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] 作者. 文献标题[J]. 期刊名, 年份, 卷(期): 页码.")

    doc.add_heading("AI 工具使用声明", level=1)
    _add_placeholder(doc, "按实际情况说明工具、用途、人工核验方式与使用日期；未使用时明确声明未使用。")

    doc.add_heading("附录：支撑文件/材料清单", level=1)
    _add_placeholder(doc, "列出程序、数据、结果文件及其用途。")
    doc.add_heading("附录：程序代码", level=1)
    _add_placeholder(doc, "附核心程序或说明支撑材料中的源文件位置。")

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


if __name__ == "__main__":
    import sys

    build_template(sys.argv[1] if len(sys.argv) > 1 else "paper_template.docx")
