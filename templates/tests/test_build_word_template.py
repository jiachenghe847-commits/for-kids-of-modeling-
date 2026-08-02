import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from build_word_template import build_template

EXPECTED_HEADINGS = [
    "摘要", "问题重述", "问题分析", "模型假设", "符号说明",
    "模型的建立与求解", "灵敏度分析", "模型评价与推广", "参考文献",
    "AI 工具使用声明", "附录：支撑文件/材料清单", "附录：程序代码",
]


def test_build_template_creates_all_headings(tmp_path):
    out = tmp_path / "paper_template.docx"
    build_template(str(out))
    doc = Document(str(out))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for expected in EXPECTED_HEADINGS:
        assert expected in headings


def test_build_template_uses_competition_page_layout(tmp_path):
    out = tmp_path / "paper_template.docx"
    build_template(str(out))
    doc = Document(str(out))
    section = doc.sections[0]

    assert section.page_width.cm == pytest.approx(21.0, abs=0.02)
    assert section.page_height.cm == pytest.approx(29.7, abs=0.02)
    assert section.left_margin.cm == pytest.approx(2.35, abs=0.02)
    assert section.right_margin.cm == pytest.approx(2.35, abs=0.02)
    assert section.top_margin.cm == pytest.approx(2.10, abs=0.02)
    assert section.bottom_margin.cm == pytest.approx(2.20, abs=0.02)


def test_build_template_defines_fonts_captions_and_page_number(tmp_path):
    out = tmp_path / "paper_template.docx"
    build_template(str(out))
    doc = Document(str(out))

    assert doc.styles["Normal"].font.size == Pt(12)
    assert doc.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert doc.styles["Heading 1"].element.rPr.rFonts.get(qn("w:eastAsia")) == "黑体"
    assert doc.styles["Title"].font.size == Pt(22)
    assert "图题" in doc.styles
    assert "表题" in doc.styles
    assert "表注" in doc.styles

    footer_xml = doc.sections[0].footer._element.xml
    assert "PAGE" in footer_xml
